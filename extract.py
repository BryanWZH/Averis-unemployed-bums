"""Extract the 7 comparison fields from an SI/BL attachment, whatever its
format (txt / pdf / docx / xlsx), and flag documents that can't be read
confidently (wrong document type, empty/corrupt file, image-only scan).
"""
import io
import re

from normalize import field_for_label, match_prefix_label, COMPARE_FIELDS

WRONG_DOC_TITLES = ["COMMERCIAL INVOICE", "PACKING LIST", "CERTIFICATE OF ORIGIN"]
GOOD_DOC_TITLES = ["SHIPPING INSTRUCTION", "BILL OF LADING", "BL INSTRUCTION", "S.I."]


class ExtractResult:
    def __init__(self):
        self.readable = True
        self.unreadable_reason = None   # detail string, for our own debugging
        self.wrong_doc_type = False
        self.title = None
        self.fields = {}          # canonical field -> raw string value
        self.ocr_used = False
        self.raw_text = ""        # best-effort flat text, for title/type checks


def _first_nonempty_line(text):
    for ln in text.splitlines():
        ln = ln.strip()
        if ln:
            return ln
    return ""


def _parse_labelvalue_lines(lines):
    """Given a list of text lines, pull out {field: value} for every line
    that starts with a known label followed by ':'."""
    fields = {}
    for ln in lines:
        if ":" not in ln:
            continue
        label, _, value = ln.partition(":")
        field = field_for_label(label)
        if field and field not in fields:
            val = value.strip()
            if val:
                fields[field] = val
    return fields


def _classify_doc_shape(raw_text, fields):
    """Decide wrong_doc_type from title markers + how many of the 7 fields
    we actually recovered."""
    title = _first_nonempty_line(raw_text).upper()
    for bad in WRONG_DOC_TITLES:
        if bad in title:
            return True, title
    if "NOT A SHIPPING INSTRUCTION" in raw_text.upper() or \
       "NOT AN SI OR BL" in raw_text.upper() or \
       "NOT THE DRAFT BL" in raw_text.upper():
        return True, title
    # generic fallback: a genuine SI/BL always yields at least the port +
    # weight/container fields; a foreign document type won't.
    core = {"port_of_loading", "port_of_discharge", "gross_weight_kg", "container_count"}
    if len(core & set(fields)) == 0 and len(fields) <= 1:
        return True, title
    return False, title


# ---------------------------------------------------------------------------
# Per-format readers. Each returns (raw_text, fields_dict) or raises.
# ---------------------------------------------------------------------------
def _read_txt(data: bytes):
    text = data.decode("utf-8", errors="replace")
    fields = _parse_labelvalue_lines(text.splitlines())
    return text, fields


def _read_xlsx(data: bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    lines = []
    fields = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        label = row[0]
        value = row[1] if len(row) > 1 else None
        if label is None:
            continue
        label_s = str(label)
        lines.append(f"{label_s}: {value}" if value is not None else label_s)
        field = field_for_label(label_s)
        if field and value is not None and field not in fields:
            fields[field] = str(value)
    return "\n".join(lines), fields


def _read_docx(data: bytes):
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    fields = {}
    for table in d.tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 2:
                continue
            label_s = cells[0].text
            value_s = cells[1].text
            lines.append(f"{label_s}: {value_s}")
            field = field_for_label(label_s)
            if field and field not in fields:
                first_line = value_s.split("\n")[0].strip()
                if first_line:
                    fields[field] = first_line
    return "\n".join(lines), fields


def _join_chars(chars):
    """Reconstruct text from a run of pdfplumber char dicts, inserting a
    space wherever there's a real gap between consecutive glyphs."""
    chars = sorted(chars, key=lambda c: c["x0"])
    out = []
    prev_x1 = None
    for c in chars:
        if prev_x1 is not None and c["x0"] - prev_x1 > 1.0:
            out.append(" ")
        out.append(c["text"])
        prev_x1 = c["x1"]
    return "".join(out).strip()


def _pdf_fields_by_font(data: bytes):
    """Label/value pairs in the PDF are drawn in different font weights
    (label = Helvetica-Bold, value = Helvetica) at fixed x-columns. A
    long label can visually run into the value column, which corrupts
    plain left-to-right text extraction (glyphs interleave once their
    x-ranges overlap). Splitting characters by font weight before
    reconstructing each side sidesteps that entirely."""
    import pdfplumber
    fields = {}
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            lines = {}
            for c in page.chars:
                lines.setdefault(round(c["top"]), []).append(c)
            for cs in lines.values():
                bold = [c for c in cs if "Bold" in c.get("fontname", "")]
                reg = [c for c in cs if "Bold" not in c.get("fontname", "")]
                if not bold or not reg:
                    continue
                label_text = _join_chars(bold)
                value_text = _join_chars(reg)
                field = field_for_label(label_text)
                if field and value_text and field not in fields:
                    fields[field] = value_text
    return fields


def _pdf_text_pdfplumber(data: bytes):
    import pdfplumber
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        parts = []
        for page in pdf.pages:
            t = page.extract_text() or ""
            parts.append(t)
        return "\n".join(parts)


def _pdf_ocr(data: bytes):
    """Rasterise each page and OCR it. Returns '' if this can't be done
    at all (e.g. the bytes aren't a real PDF)."""
    try:
        from pdf2image import convert_from_bytes
        import pytesseract
        images = convert_from_bytes(data, dpi=200)
        parts = []
        for img in images:
            parts.append(pytesseract.image_to_string(img))
        return "\n".join(parts)
    except Exception:
        return ""


def _parse_pdf_lines(text):
    """PDF label/value pairs sit on the SAME text line (label drawn at one
    x position, value at another), and pdfplumber's plain-text extraction
    can collapse that column gap down to a single space -- so a colon
    isn't a safe assumption. Try, in order: 'Label: value', a known label
    as a literal line prefix (handles the single-space-gap case and the
    'TOTAL <label>: value' weight line), then a wider whitespace gap."""
    fields = {}
    for raw_ln in text.splitlines():
        ln = raw_ln.strip()
        if not ln:
            continue
        if ":" in ln:
            label, _, value = ln.partition(":")
            field = field_for_label(label)
            if field and value.strip() and field not in fields:
                fields[field] = value.strip()
                continue
        hit = match_prefix_label(ln)
        if hit:
            field, value = hit
            if field not in fields:
                fields[field] = value
            continue
        # column layout with 2+ spaces between label and value
        m = re.match(r"^(.{2,40}?)\s{2,}(.+)$", ln)
        if m:
            field = field_for_label(m.group(1))
            if field and field not in fields:
                fields[field] = m.group(2).strip()
    return fields


def extract(data: bytes, filename: str) -> ExtractResult:
    res = ExtractResult()
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if len(data) == 0:
        res.readable = False
        res.unreadable_reason = "empty_file"
        return res

    try:
        if ext == "txt":
            text, fields = _read_txt(data)
        elif ext == "xlsx":
            text, fields = _read_xlsx(data)
        elif ext == "docx":
            text, fields = _read_docx(data)
        elif ext == "pdf":
            text = _pdf_text_pdfplumber(data)
            fields = _parse_pdf_lines(text)
            try:
                font_fields = _pdf_fields_by_font(data)
                fields.update(font_fields)  # font-split reading wins: it
                                             # can't be corrupted by label/
                                             # value column overlap
            except Exception:
                pass
            if len(re.sub(r"[^A-Za-z0-9]", "", text)) < 25:
                # no usable text layer -> try OCR, but still flag for
                # human confirmation (scanned quantities/names are exactly
                # the kind of thing worth a second pair of eyes on).
                ocr_text = _pdf_ocr(data)
                res.ocr_used = True
                if len(re.sub(r"[^A-Za-z0-9]", "", ocr_text)) < 25:
                    res.readable = False
                    res.unreadable_reason = "no_text_layer_and_ocr_failed"
                    return res
                text = ocr_text
                fields = _parse_pdf_lines(text)
                res.readable = False
                res.unreadable_reason = "image_only_pdf_ocr_used"
                res.raw_text = text
                res.fields = fields
                return res
        else:
            res.readable = False
            res.unreadable_reason = f"unsupported_extension:{ext}"
            return res
    except Exception as exc:
        res.readable = False
        res.unreadable_reason = f"parse_error:{type(exc).__name__}"
        return res

    res.raw_text = text
    res.fields = fields
    wrong, title = _classify_doc_shape(text, fields)
    res.title = title
    res.wrong_doc_type = wrong
    return res
